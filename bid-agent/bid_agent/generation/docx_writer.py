"""将标书导出为 Word（.docx）。支持整本导出与分项导出。"""
from __future__ import annotations

from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..models import BidDocument, BidSection, RequirementSummary


def _set_run_cjk(run, font_name: str) -> None:
    """为单个 run 设置中文（eastAsia）字体，确保 rPr/rFonts 存在。"""
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def _set_cjk_font(doc: Document, font_name: str = "宋体", size: int = 12) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_name)


def _heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    _set_run_cjk(run, "黑体")
    run.font.color.rgb = RGBColor(0, 0, 0)


def _para(doc: Document, text: str) -> None:
    for block in text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(24)
        p.add_run(block)


def _table(doc: Document, rows: List[List[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(cols):
            val = row[ci] if ci < len(row) else ""
            cells[ci].text = str(val)
            if ri == 0:
                for p in cells[ci].paragraphs:
                    for run in p.runs:
                        run.font.bold = True


def write_bid_docx(
    bid: BidDocument,
    out_path: Path,
    req: Optional[RequirementSummary] = None,
) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _set_cjk_font(doc)

    # ---- 封面 ----
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(bid.project_name or "投标文件")
    run.font.size = Pt(28)
    run.font.bold = True
    _set_run_cjk(run, "黑体")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("投 标 文 件")
    r2.font.size = Pt(22)
    _set_run_cjk(r2, "黑体")

    for _ in range(6):
        doc.add_paragraph()
    for label, value in [
        ("投标人（盖章）", bid.company or "____________"),
        ("项目编号", (req.project_no if req else "") or "____________"),
        ("采购人", (req.purchaser if req else "") or "____________"),
        ("日期", "______年____月____日"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(f"{label}：{value}")
        rr.font.size = Pt(14)
        _set_run_cjk(rr, "宋体")
    doc.add_page_break()

    # ---- 目录（占位说明）----
    _heading(doc, "目  录", level=1)
    for i, s in enumerate(bid.sections, 1):
        doc.add_paragraph(f"{i}. {s.title}")
    doc.add_page_break()

    # ---- 正文 ----
    for i, s in enumerate(bid.sections, 1):
        _heading(doc, f"{i}、{s.title}", level=1)
        if s.content:
            _para(doc, s.content)
        if s.table:
            doc.add_paragraph()
            _table(doc, s.table)
        if s.sources:
            note = doc.add_paragraph()
            nr = note.add_run("【参考企业资料：" + "、".join(s.sources) + "】")
            nr.font.size = Pt(9)
            nr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        doc.add_page_break()

    doc.save(str(out_path))
    return out_path


def write_section_docx(section: BidSection, out_path: Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _set_cjk_font(doc)
    _heading(doc, section.title, level=1)
    if section.content:
        _para(doc, section.content)
    if section.table:
        _table(doc, section.table)
    doc.save(str(out_path))
    return out_path
