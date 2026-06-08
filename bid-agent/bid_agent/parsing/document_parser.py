"""招标文件解析：提取完整文本、表格与结构（标题层级、页码）。

  - Word（.docx）：python-docx，按正文顺序解析段落与表格，依样式识别标题。
  - PDF（.pdf）：pypdf 提取逐页文本；pdfplumber 提取表格（若可用）。
  - 文本（.txt/.md）：直接读取。
"""
from __future__ import annotations

from pathlib import Path
from typing import List

from ..logging_utils import get_logger
from ..models import DocElement, ParsedDocument
from .cleaner import clean_text, merge_broken_lines, remove_repeated_headers

log = get_logger("parsing")

SUPPORTED_EXTS = {".docx", ".pdf", ".txt", ".md", ".markdown"}


def parse_document(path: Path) -> ParsedDocument:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".docx":
        return _parse_docx(path)
    if ext == ".pdf":
        return _parse_pdf(path)
    if ext in {".txt", ".md", ".markdown"}:
        return _parse_text(path)
    if ext == ".doc":
        raise ValueError("不支持旧版 .doc，请另存为 .docx 后上传。")
    raise ValueError(f"不支持的文件类型：{ext}")


def _table_to_rows(rows) -> List[List[str]]:
    return [[clean_text(c) for c in r] for r in rows]


def _parse_docx(path: Path) -> ParsedDocument:
    import docx
    from docx.document import Document as _Doc
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    elements: List[DocElement] = []
    text_parts: List[str] = []
    table_count = 0

    body = doc.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            para = Paragraph(child, doc)
            txt = clean_text(para.text)
            if not txt:
                continue
            style = (para.style.name or "").lower() if para.style else ""
            if style.startswith("heading") or "标题" in (para.style.name or ""):
                level = _heading_level(para.style.name)
                elements.append(DocElement(kind="heading", text=txt, level=level))
            else:
                elements.append(DocElement(kind="paragraph", text=txt))
            text_parts.append(txt)
        elif isinstance(child, CT_Tbl):
            table = Table(child, doc)
            rows = _table_to_rows([[cell.text for cell in row.cells] for row in table.rows])
            elements.append(DocElement(kind="table", rows=rows))
            table_count += 1
            text_parts.append(_render_table_text(rows))

    full_text = clean_text("\n".join(text_parts))
    return ParsedDocument(
        filename=path.name,
        file_type="docx",
        full_text=full_text,
        elements=elements,
        page_count=_estimate_pages(full_text),
        char_count=len(full_text),
        table_count=table_count,
    )


def _heading_level(style_name: str) -> int:
    import re

    m = re.search(r"(\d+)", style_name or "")
    return int(m.group(1)) if m else 1


def _render_table_text(rows: List[List[str]]) -> str:
    return "\n".join(" | ".join(r) for r in rows if any(c.strip() for c in r))


def _parse_pdf(path: Path) -> ParsedDocument:
    elements: List[DocElement] = []
    warnings: List[str] = []
    page_texts: List[str] = []
    table_count = 0

    # 文本：pypdf
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        for page in reader.pages:
            page_texts.append(page.extract_text() or "")
    except Exception as e:
        warnings.append(f"pypdf 文本提取失败：{e}")

    page_texts = remove_repeated_headers(page_texts)

    # 表格：pdfplumber（可选）
    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            for pno, page in enumerate(pdf.pages, start=1):
                for tbl in page.extract_tables() or []:
                    rows = _table_to_rows([[c or "" for c in row] for row in tbl])
                    if rows:
                        elements.append(DocElement(kind="table", rows=rows, page=pno))
                        table_count += 1
    except Exception as e:
        warnings.append(f"pdfplumber 表格提取跳过：{e}")

    for pno, ptxt in enumerate(page_texts, start=1):
        ptxt = clean_text(merge_broken_lines(ptxt))
        if ptxt:
            elements.append(DocElement(kind="paragraph", text=ptxt, page=pno))

    full_text = clean_text("\n".join(p for p in page_texts if p))
    return ParsedDocument(
        filename=path.name,
        file_type="pdf",
        full_text=full_text,
        elements=elements,
        page_count=len(page_texts),
        char_count=len(full_text),
        table_count=table_count,
        warnings=warnings,
    )


def _parse_text(path: Path) -> ParsedDocument:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    full_text = clean_text(raw)
    elements = [DocElement(kind="paragraph", text=p) for p in full_text.split("\n\n") if p.strip()]
    return ParsedDocument(
        filename=path.name,
        file_type=path.suffix.lstrip("."),
        full_text=full_text,
        elements=elements,
        page_count=_estimate_pages(full_text),
        char_count=len(full_text),
    )


def _estimate_pages(text: str) -> int:
    # 粗略：每页约 1200 字
    return max(1, round(len(text) / 1200))
