# -*- coding: utf-8 -*-
"""客户《安规考试题库 0407 更新》→ bank.js（QBANK）；制度文件夹内两个题库 → ZBANK。
用法：python3 gen_bank.py   （读取 materials/1002客户修改意见/，输出 bank.js；build.py 拼接）
每个题库：参考文件 / 单选题 / 多选题 / 判断题 / 小案例题 / 大案例题。
分类（雷达维度，按内容分类，docs/02 第 69 条 ③）在这里按关键词打标，页面不再做正则。"""
import os, re, json, glob, sys
import openpyxl
try:
    import xlrd
except Exception:
    xlrd = None

B = os.path.dirname(os.path.abspath(__file__))
TK = os.path.join(B, 'materials', '1002客户修改意见', '安规考试题库0407更新')
ZD = os.path.join(B, 'materials', '1002客户修改意见', '制度文件', '2-事故隐患判定标准')
OUT = os.path.join(B, 'bank.js')

# ---------- 维度分类（关键词按优先级匹配，首个命中即归类） ----------
RULE_CATS = [
    ('r_emerg', '应急与消防', r'急救|心肺|消防|灭火|火灾|应急|抢修|触电者|脱离电源'),
    ('r_special', '带电与特殊作业', r'带电作业|高处|坠落|起重|吊装|吊车|动火|有限空间|试验|二次|保护装置|压板|电缆|焊|登高|绝缘斗|绝缘杆'),
    ('r_tech', '保证安全的技术措施', r'停电|验电|接地|遮栏|围栏|标示牌|标志牌|接地线|接地刀闸|验电器|明显断开'),
    ('r_switch', '倒闸操作与防误', r'操作票|倒闸|操作人|监护操作|五防|防误|闭锁|解锁|双重名称|隔离开关|断路器|拉合|合上|拉开|操作任务|模拟'),
    ('r_ticket', '工作票与许可', r'工作票|许可|工作负责人|签发|监护人|工作班|间断|转移|终结|延期|派工单|勘察|交底'),
    ('r_base', '基本要求与人员条件', r'.'),
]
LIFE_CATS = [
    ('l_aid', '应急与急救', r'急救|心肺|触电急救|灭火|火灾|应急|撤离|脱离电源'),
    ('l_high', '高处起重与特殊作业', r'高处|坠落|起重|吊|脚手|梯|有限空间|动火|气体|带电作业'),
    ('l_ppe', '个人防护与工器具', r'安全帽|安全带|绝缘手套|绝缘靴|护目|工器具|验电器|试验合格|防护用品|劳动防护'),
    ('l_tech', '停电验电接地', r'停电|验电|接地|遮栏|标示牌|标志牌|围栏|明显断开'),
    ('l_perm', '作业许可与监护', r'工作票|许可|监护|负责人|签发|擅自|单独|无票|操作票|交底'),
    ('l_dist', '安全距离与防触电', r'.'),
]
CASE_CATS = [
    ('c_cause', '原因分析', r'原因'),
    ('c_viol', '违规辨识', r'违反|违章|违规|禁止|严禁|五防|不得|不应'),
    ('c_measure', '防范措施与教训', r'措施|防范|避免|教训|如何|改进|整改|防止'),
    ('c_duty', '责任与职责', r'责任|职责|负责人|许可人|监护人|签发人|操作人|应由'),
    ('c_hazard', '事故定性与隐患识别', r'隐患|事故等级|定性|属于|事件|级别|分级'),
    ('c_rule', '规程依据', r'.'),
]
def cat_of(cats, text):
    for k, n, re_ in cats:
        if re.search(re_, text): return k
    return cats[-1][0]

# ---------- 读取 ----------
def rows_of(path, sheet):
    if path.endswith('.xls'):
        if not xlrd: return []
        wb = xlrd.open_workbook(path)
        if sheet not in wb.sheet_names(): return []
        ws = wb.sheet_by_name(sheet)
        return [[ws.cell_value(r, c) for c in range(ws.ncols)] for r in range(ws.nrows)]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames: return []
    return [list(r) for r in wb[sheet].iter_rows(values_only=True)]
def sheets_of(path):
    if path.endswith('.xls'):
        return xlrd.open_workbook(path).sheet_names() if xlrd else []
    return openpyxl.load_workbook(path, read_only=True).sheetnames
def s(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).replace('　', ' ').strip()
def header_idx(rows):
    for i, r in enumerate(rows[:6]):
        cells = [s(c) for c in r]
        if '题干' in cells: return i, cells
    return None, None
def col(hdr, *names):
    for n in names:
        for i, c in enumerate(hdr):
            if c and n in c: return i
    return -1

def parse_objective(rows, kind):
    hi, hdr = header_idx(rows)
    if hi is None: return []
    iq, ia = col(hdr, '题干'), col(hdr, '答案')
    io = [i for i, c in enumerate(hdr) if c in list('ABCDEFGHI')]
    ib, ik, isub, il, idf, icat = col(hdr, '制度名称'), col(hdr, '关键词'), col(hdr, '专业小类'), col(hdr, '保命'), col(hdr, '难易'), col(hdr, '考试类别', '知识类别')
    out = []
    for r in rows[hi + 1:]:
        r = list(r) + [None] * 24
        stem = s(r[iq]); ans = s(r[ia]).upper().replace(' ', '')
        if not stem or not ans or not s(r[0]).isdigit(): continue
        opts = [s(r[i]) for i in io]
        while opts and not opts[-1]: opts.pop()
        q = {'t': kind, 'stem': stem, 'ans': ans}
        if kind == 'judge':
            q['ans'] = ans in ('A', '正确', '对', 'T', 'TRUE', '√')
            q['opts'] = []
        else:
            q['opts'] = opts
            if kind == 'single' and len(ans) != 1: continue
        q['basis'] = s(r[ib]) if ib >= 0 else ''
        kw = s(r[ik]) if ik >= 0 else ''
        q['life'] = (s(r[il]) == '是') if il >= 0 else False
        q['cat'] = cat_of(RULE_CATS, stem + ' ' + q['basis'] + ' ' + kw)
        if q['life']: q['lcat'] = cat_of(LIFE_CATS, stem + ' ' + q['basis'] + ' ' + kw)
        out.append(q)
    return out

def parse_cases(rows, kind):
    hi, hdr = header_idx(rows)
    if hi is None: return []
    iq, ia, ino = col(hdr, '题干'), col(hdr, '答案'), col(hdr, '题号')
    io = [i for i, c in enumerate(hdr) if c in list('ABCDEFGHI')]
    ib, ik, idf = col(hdr, '制度名称'), col(hdr, '关键词'), col(hdr, '难易')
    cases, cur = [], None
    for r in rows[hi + 1:]:
        r = list(r) + [None] * 24
        no, stem = s(r[ino]) if ino >= 0 else '', s(r[iq])
        if not stem: continue
        if no == '题干' or (not cur and not s(r[ia])):
            cur = {'kind': kind, 'stem': stem, 'qs': []}; cases.append(cur); continue
        if not cur: continue
        ans = s(r[ia]).upper().replace(' ', '')
        if not ans: continue
        opts = [s(r[i]) for i in io]
        while opts and not opts[-1]: opts.pop()
        judge = ('判断' in stem) or (len(opts) == 2 and opts[0] in ('正确', '对') and opts[1] in ('错误', '错'))
        q = {'stem': stem, 'basis': s(r[ib]) if ib >= 0 else ''}
        if judge:
            q['t'] = 'judge'; q['ans'] = ans in ('A', '正确', '对', 'T', 'TRUE', '√'); q['opts'] = []
        else:
            q['t'] = 'multi' if len(ans) > 1 or '多选' in stem else 'single'; q['ans'] = ans; q['opts'] = opts
        q['cat'] = cat_of(CASE_CATS, stem)
        cur['qs'].append(q)
    return [c for c in cases if c['qs']]

def parse_bank(path):
    sh = sheets_of(path)
    refs = []
    if any('参考文件' in x for x in sh):
        for r in rows_of(path, next(x for x in sh if '参考文件' in x)):
            r = [s(c) for c in r if s(c)]
            if len(r) >= 2 and r[0].isdigit(): refs.append(r[1])
    pick = lambda name: next((x for x in sh if name in x), None)
    q = []
    for name, kind in [('单选题', 'single'), ('多选题', 'multi'), ('判断题', 'judge')]:
        if pick(name): q += parse_objective(rows_of(path, pick(name)), kind)
    cases = []
    for name, kind in [('小案例题', 'small'), ('大案例题', 'big')]:
        if pick(name): cases += parse_cases(rows_of(path, pick(name)), kind)
    return refs, q, cases

def main():
    pros = []
    for f in sorted(glob.glob(os.path.join(TK, '**', '*.xls*'), recursive=True)):
        grp = os.path.basename(os.path.dirname(f))
        m = re.search(r'题库-(.+?)\.(xlsx|xls)$', os.path.basename(f))
        name = m.group(1) if m else os.path.basename(f)
        grp = re.sub(r'^\d+\.', '', grp)
        try:
            refs, q, cases = parse_bank(f)
        except Exception as e:
            print('skip', f, e); continue
        if not q: print('empty', f); continue
        pid = 'p' + str(len(pros) + 1)
        pros.append({'id': pid, 'n': name, 'g': grp, 'f': os.path.basename(f), 'refs': refs, 'q': q, 'cases': cases})
        print('%-14s %-22s 题 %4d（保命 %3d）案例 %2d  单%d 多%d 判%d' % (grp, name, len(q), sum(1 for x in q if x['life']), len(cases), sum(1 for x in q if x['t'] == 'single'), sum(1 for x in q if x['t'] == 'multi'), sum(1 for x in q if x['t'] == 'judge')))
    # 默认专业：变电运行类
    default = next((p['id'] for p in pros if '变电运行' in p['n']), pros[0]['id'])
    # 制度题库（事故隐患 / 硬措施）
    zb = []
    for f in sorted(glob.glob(os.path.join(ZD, '*.xlsx'))):
        for sh in sheets_of(f):
            rows = rows_of(f, sh)
            hi, hdr = header_idx(rows)
            if hi is None: continue
            iq, ia, it, ib = col(hdr, '题干'), col(hdr, '答案'), col(hdr, '题型'), col(hdr, '出自规范', '制度名称')
            io = [i for i, c in enumerate(hdr) if c in list('ABCDEFGHI')]
            for r in rows[hi + 1:]:
                r = list(r) + [None] * 20
                stem, ans, typ = s(r[iq]), s(r[ia]).upper().replace(' ', ''), s(r[it])
                if not stem or not ans: continue
                opts = [s(r[i]) for i in io]
                while opts and not opts[-1]: opts.pop()
                if '判断' in typ: q = {'t': 'judge', 'stem': stem, 'ans': ans in ('A', '正确', '对', 'T', 'TRUE', '√'), 'opts': []}
                elif '多选' in typ or len(ans) > 1: q = {'t': 'multi', 'stem': stem, 'ans': ans, 'opts': opts}
                elif '填空' in typ or '简答' in typ: continue
                else: q = {'t': 'single', 'stem': stem, 'ans': ans, 'opts': opts}
                if q['t'] != 'judge' and not q['opts']: continue
                q['basis'] = s(r[ib]) if ib >= 0 else ''
                q['src'] = '41号令测试题库' if '41' in os.path.basename(f) else '事故隐患与安全生产硬措施考试复习资料'
                q['theme'] = 'z3' if re.search(r'硬措施|十条', stem + q['basis']) else 'z2'
                zb.append(q)
    print('制度题库', len(zb), '题；z2', sum(1 for x in zb if x['theme'] == 'z2'), 'z3', sum(1 for x in zb if x['theme'] == 'z3'))
    cats = {'rule': [[k, n] for k, n, _ in RULE_CATS], 'life': [[k, n] for k, n, _ in LIFE_CATS], 'case': [[k, n] for k, n, _ in CASE_CATS]}
    js = '/* 由 gen_bank.py 从客户《安规考试题库 0407 更新》与制度题库生成，勿手改 */\n'
    js += 'const QBANK = ' + json.dumps({'pros': pros, 'default': default, 'cats': cats}, ensure_ascii=False, separators=(',', ':')) + ';\n'
    js += 'const ZBANK = ' + json.dumps(zb, ensure_ascii=False, separators=(',', ':')) + ';\n'
    open(OUT, 'w', encoding='utf-8').write(js)
    # 分布
    from collections import Counter
    for p in pros:
        if p['id'] == default:
            print('默认', p['n'], 'rule cats', Counter(x['cat'] for x in p['q']).most_common(), '| life cats', Counter(x['lcat'] for x in p['q'] if x['life']).most_common(), '| case cats', Counter(q['cat'] for c in p['cases'] for q in c['qs']).most_common())
    print('bank.js', len(js), 'bytes')

if __name__ == '__main__':
    main()
